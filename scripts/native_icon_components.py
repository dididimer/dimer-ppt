#!/usr/bin/env python3
"""Replace selected raster icons with small, editable PowerPoint shape families.

The map is explicit: automatic image recognition cannot reliably distinguish a
flat icon from a photograph or text screenshot. Each replacement uses the
picture's existing box and z-order. The source deck is never modified.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Pt


COLORS = {
    "ink": "263340",
    "outline": "263340",
    "blue": "7FB6DB",
    "pale_blue": "D8EAF6",
    "face": "A5CEE8",
    "white": "FFFFFF",
    "gold": "F5BF27",
    "gold_light": "FFE47A",
    "gold_dark": "C88B12",
    "red": "E96F72",
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


class IconCanvas:
    def __init__(self, slide, x: int, y: int, w: int, h: int, design_w: int, design_h: int, prefix: str):
        self.slide = slide
        self.x, self.y, self.w, self.h = x, y, w, h
        self.design_w, self.design_h = design_w, design_h
        self.prefix = prefix
        self.created = []

    def px(self, value: float) -> int:
        return self.x + round(value * self.w / self.design_w)

    def py(self, value: float) -> int:
        return self.y + round(value * self.h / self.design_h)

    def sw(self, value: float) -> int:
        return max(1, round(value * self.w / self.design_w))

    def sh(self, value: float) -> int:
        return max(1, round(value * self.h / self.design_h))

    def _style(self, shape, label: str, fill: str | None, stroke: str | None, width_pt: float = 0.9):
        shape.name = f"{self.prefix}_{label}_{len(self.created):02d}"
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(COLORS.get(fill, fill))
        else:
            shape.fill.background()
        if stroke:
            shape.line.color.rgb = rgb(COLORS.get(stroke, stroke))
            shape.line.width = Pt(width_pt)
        else:
            shape.line.fill.background()
        for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
            effect_ref.set("idx", "0")
        self.created.append(shape)
        return shape

    def box(self, label: str, kind, x: float, y: float, w: float, h: float, fill: str, stroke: str | None = None, width_pt: float = 0.9):
        shape = self.slide.shapes.add_shape(kind, self.px(x), self.py(y), self.sw(w), self.sh(h))
        return self._style(shape, label, fill, stroke, width_pt)

    def line(self, label: str, x1: float, y1: float, x2: float, y2: float, stroke: str = "ink", width_pt: float = 0.9):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.px(x1), self.py(y1), self.px(x2), self.py(y2))
        shape.name = f"{self.prefix}_{label}_{len(self.created):02d}"
        shape.line.color.rgb = rgb(COLORS.get(stroke, stroke))
        shape.line.width = Pt(width_pt)
        for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
            effect_ref.set("idx", "0")
        self.created.append(shape)
        return shape

    def polygon(self, label: str, points: list[tuple[float, float]], fill: str, stroke: str | None = None, width_pt: float = 0.9):
        if len(points) < 3:
            raise ValueError("polygon needs at least three points")
        builder = self.slide.shapes.build_freeform(points[0][0], points[0][1], scale=(self.w / self.design_w, self.h / self.design_h))
        builder.add_line_segments(points[1:], close=True)
        shape = builder.convert_to_shape(origin_x=self.x, origin_y=self.y)
        return self._style(shape, label, fill, stroke, width_pt)


def draw_robot(c: IconCanvas) -> None:
    # Familiar PRM robot: antenna, ears, rounded head, visor, eyes, neck.
    c.line("antenna", 29, 2, 29, 11, width_pt=0.8)
    c.box("antenna_tip", MSO_SHAPE.OVAL, 26, 0, 6, 6, "blue", "ink", 0.7)
    c.box("left_ear", MSO_SHAPE.ROUNDED_RECTANGLE, 3, 20, 9, 18, "blue", "ink")
    c.box("right_ear", MSO_SHAPE.ROUNDED_RECTANGLE, 46, 20, 9, 18, "blue", "ink")
    c.box("head", MSO_SHAPE.ROUNDED_RECTANGLE, 9, 10, 40, 34, "white", "ink", 1.1)
    c.box("forehead", MSO_SHAPE.ROUNDED_RECTANGLE, 15, 13, 28, 8, "pale_blue")
    c.box("visor", MSO_SHAPE.ROUNDED_RECTANGLE, 13, 21, 32, 18, "face", "ink", 0.8)
    c.box("left_eye", MSO_SHAPE.OVAL, 20, 27, 5, 5, "ink")
    c.box("right_eye", MSO_SHAPE.OVAL, 33, 27, 5, 5, "ink")
    c.box("left_glint", MSO_SHAPE.OVAL, 21, 27, 2, 2, "white")
    c.box("right_glint", MSO_SHAPE.OVAL, 34, 27, 2, 2, "white")
    c.box("neck", MSO_SHAPE.RECTANGLE, 24, 44, 10, 4, "pale_blue", "ink", 0.7)
    c.box("base", MSO_SHAPE.ROUNDED_RECTANGLE, 17, 48, 24, 6, "pale_blue", "ink", 0.7)


def draw_shield(c: IconCanvas) -> None:
    c.polygon("shield", [(21, 2), (37, 8), (36, 27), (31, 37), (21, 44), (11, 37), (6, 27), (5, 8)], "blue", "ink", 1.0)
    c.polygon("check", [(12, 22), (17, 27), (29, 15), (32, 18), (17, 32), (9, 25)], "white")
    c.line("highlight", 10, 11, 18, 7, "white", 0.7)


def draw_trophy(c: IconCanvas) -> None:
    c.box("left_handle", MSO_SHAPE.OVAL, 3, 9, 18, 23, "white", "ink", 1.0)
    c.box("right_handle", MSO_SHAPE.OVAL, 55, 9, 18, 23, "white", "ink", 1.0)
    c.polygon("cup", [(16, 6), (60, 6), (57, 26), (52, 35), (44, 40), (32, 40), (23, 35), (19, 26)], "gold", "ink", 1.0)
    c.box("rim", MSO_SHAPE.ROUNDED_RECTANGLE, 15, 4, 46, 7, "gold_light", "ink", 0.8)
    c.box("stem", MSO_SHAPE.RECTANGLE, 33, 39, 10, 13, "gold", "ink", 0.8)
    c.box("foot", MSO_SHAPE.ROUNDED_RECTANGLE, 22, 51, 32, 6, "gold_light", "ink", 0.8)
    c.box("base", MSO_SHAPE.ROUNDED_RECTANGLE, 16, 57, 44, 9, "gold", "ink", 0.9)
    c.polygon("star", [(38, 14), (40, 19), (45, 19), (41, 23), (43, 28), (38, 25), (33, 28), (35, 23), (31, 19), (36, 19)], "gold_dark")


def paper(c: IconCanvas, label: str, x: float, y: float, w: float, h: float, *, rows: int = 3) -> None:
    fold = min(w, h) * 0.22
    c.polygon(label, [(x, y), (x + w - fold, y), (x + w, y + fold), (x + w, y + h), (x, y + h)], "white", "ink", 0.9)
    c.polygon(f"{label}_fold", [(x + w - fold, y), (x + w - fold, y + fold), (x + w, y + fold)], "pale_blue", "ink", 0.7)
    for index in range(rows):
        yy = y + h * (0.43 + index * 0.15)
        c.line(f"{label}_rule", x + w * 0.19, yy, x + w * 0.78, yy, "blue", 0.7)


def check_badge(c: IconCanvas, x: float, y: float, d: float) -> None:
    c.box("check_badge", MSO_SHAPE.OVAL, x, y, d, d, "68BE7A", "ink", 0.8)
    c.line("check_short", x + d * 0.20, y + d * 0.52, x + d * 0.41, y + d * 0.72, "white", 1.5)
    c.line("check_long", x + d * 0.41, y + d * 0.72, x + d * 0.82, y + d * 0.27, "white", 1.5)


def draw_top_dialogue(c: IconCanvas) -> None:
    paper(c, "note", 65, 9, 64, 72, rows=4)
    c.box("bubble", MSO_SHAPE.ROUNDED_RECTANGLE, 5, 24, 59, 38, "face", "ink", 0.9)
    c.polygon("bubble_tail", [(28, 59), (36, 70), (41, 59)], "face", "ink", 0.7)
    for xx in (22, 35, 48):
        c.box("bubble_dot", MSO_SHAPE.OVAL, xx, 41, 5, 5, "ink")
    c.box("medical_h", MSO_SHAPE.RECTANGLE, 84, 30, 22, 6, "red")
    c.box("medical_v", MSO_SHAPE.RECTANGLE, 92, 22, 6, 22, "red")


def draw_top_candidates(c: IconCanvas) -> None:
    paper(c, "back_note", 36, 6, 70, 69, rows=2)
    paper(c, "middle_note", 24, 14, 70, 69, rows=3)
    paper(c, "front_note", 9, 22, 70, 66, rows=4)


def draw_top_gold(c: IconCanvas) -> None:
    paper(c, "gold_note", 13, 8, 61, 75, rows=4)
    check_badge(c, 50, 50, 34)


def draw_best_doc_check(c: IconCanvas) -> None:
    paper(c, "best_note", 8, 3, 63, 66, rows=3)
    check_badge(c, 51, 37, 37)


def draw_puzzle(c: IconCanvas) -> None:
    c.polygon("puzzle", [(4, 14), (13, 14), (13, 10), (15, 6), (20, 5), (24, 8), (24, 14), (31, 14), (31, 21), (36, 21), (39, 24), (37, 29), (31, 29), (31, 40), (23, 40), (23, 35), (19, 32), (15, 35), (15, 40), (4, 40), (4, 29), (9, 29), (12, 25), (9, 21), (4, 21)], "blue", "ink", 0.9)


def draw_logic(c: IconCanvas) -> None:
    c.line("left_branch", 24, 14, 11, 37, "ink", 0.8)
    c.line("right_branch", 24, 14, 39, 37, "ink", 0.8)
    c.box("root", MSO_SHAPE.OVAL, 18, 3, 13, 13, "8CCFBA", "ink", 0.8)
    c.box("left_node", MSO_SHAPE.OVAL, 5, 34, 12, 12, "face", "ink", 0.8)
    c.box("right_node", MSO_SHAPE.OVAL, 33, 34, 12, 12, "pale_blue", "ink", 0.8)


def draw_heart(c: IconCanvas) -> None:
    c.box("heart", MSO_SHAPE.HEART, 4, 7, 37, 35, "red", "ink", 0.9)
    c.box("cross_h", MSO_SHAPE.RECTANGLE, 12, 20, 21, 5, "white")
    c.box("cross_v", MSO_SHAPE.RECTANGLE, 20, 12, 5, 21, "white")


def draw_warning(c: IconCanvas) -> None:
    c.polygon("triangle", [(29, 6), (54, 52), (4, 52)], "red", "ink", 1.0)
    c.box("mark", MSO_SHAPE.ROUNDED_RECTANGLE, 26, 21, 6, 18, "white")
    c.box("dot", MSO_SHAPE.OVAL, 26, 43, 6, 6, "white")


DRAWERS = {
    "robot": (58, 56, draw_robot),
    "shield": (42, 46, draw_shield),
    "trophy": (76, 72, draw_trophy),
    "top_dialogue": (140, 90, draw_top_dialogue),
    "top_candidates": (132, 90, draw_top_candidates),
    "top_gold": (92, 94, draw_top_gold),
    "best_doc_check": (94, 72, draw_best_doc_check),
    "puzzle": (42, 46, draw_puzzle),
    "logic": (48, 50, draw_logic),
    "heart": (48, 46, draw_heart),
    "warning": (58, 60, draw_warning),
}


def replace_picture(slide, picture, icon: str):
    design_w, design_h, draw = DRAWERS[icon]
    canvas = IconCanvas(slide, picture.left, picture.top, picture.width, picture.height, design_w, design_h, f"DIMER_ICON_{icon.upper()}")
    draw(canvas)
    anchor = picture._element
    for shape in canvas.created:
        anchor.addprevious(shape._element)
    anchor.getparent().remove(anchor)
    return canvas.created


def load_map(path: Path) -> dict[int, dict[str, str]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("map must be an object keyed by 1-based slide number")
    result = {}
    for slide_number, mapping in data.items():
        if not isinstance(mapping, dict) or not all(isinstance(k, str) and v in DRAWERS for k, v in mapping.items()):
            raise ValueError(f"invalid map for slide {slide_number}; icons: {', '.join(DRAWERS)}")
        result[int(slide_number)] = mapping
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--map", required=True, type=Path, help="JSON: {\"1\": {\"Picture Name\": \"robot\"}}; supported icons: " + ", ".join(DRAWERS))
    parser.add_argument("--report", type=Path, help="Optional JSON editability report")
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        parser.error("--output must differ from --input")

    mapping = load_map(args.map)
    deck = Presentation(args.input)
    report = {"input": str(args.input), "output": str(args.output), "replaced": [], "skipped": [], "missing": [], "pictures_before": 0, "pictures_after": 0}
    for number, slide in enumerate(deck.slides, start=1):
        report["pictures_before"] += sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        requested = mapping.get(number, {})
        found = set()
        for shape in list(slide.shapes):
            if shape.name not in requested:
                continue
            found.add(shape.name)
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                raise ValueError(f"slide {number} {shape.name} is not a picture")
            if any(abs(value) > 1e-6 for value in (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom)):
                report["skipped"].append({"slide": number, "picture": shape.name, "reason": "cropped_picture"})
                continue
            transform = shape._element.spPr.xfrm
            if abs(shape.rotation) > 1e-6 or (transform is not None and (transform.flipH or transform.flipV)):
                report["skipped"].append({"slide": number, "picture": shape.name, "reason": "transformed_picture"})
                continue
            new_shapes = replace_picture(slide, shape, requested[shape.name])
            report["replaced"].append({"slide": number, "picture": shape.name, "icon": requested[shape.name], "native_shapes": len(new_shapes)})
        report["missing"].extend({"slide": number, "picture": name} for name in requested.keys() - found)
        report["pictures_after"] += sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    for number in mapping.keys() - set(range(1, len(deck.slides) + 1)):
        report["missing"].extend({"slide": number, "picture": name} for name in mapping[number])
    if report["missing"]:
        raise ValueError(f"unmatched picture names: {report['missing']}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(args.output)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
