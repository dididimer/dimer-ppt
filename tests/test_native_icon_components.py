import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "native_icon_components.py"


class NativeIconComponentsTests(unittest.TestCase):
    def make_deck(self, directory: Path, *, crop: bool = False):
        icon = directory / "icon.png"
        Image.new("RGB", (58, 56), "white").save(icon)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        before = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.2), Inches(0.3), Inches(0.3))
        before.name = "Before"
        picture = slide.shapes.add_picture(str(icon), Inches(1), Inches(1), Inches(0.6), Inches(0.6))
        picture.name = "Robot Pic"
        if crop:
            picture.crop_left = 0.1
        after = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(0.3), Inches(0.3))
        after.name = "After"
        path = directory / "original.pptx"
        deck.save(path)
        return path

    def run_replace(self, directory: Path, source: Path):
        mapping = directory / "map.json"
        mapping.write_text(json.dumps({"1": {"Robot Pic": "robot"}}), encoding="utf-8")
        output = directory / "edited.pptx"
        report = directory / "report.json"
        completed = subprocess.run(
            [sys.executable, str(SCRIPT), "--input", str(source), "--output", str(output), "--map", str(mapping), "--report", str(report)],
            capture_output=True,
            text=True,
        )
        if completed.returncode:
            self.fail(completed.stderr)
        return output, json.loads(report.read_text(encoding="utf-8"))

    def test_robot_replaces_picture_with_individually_editable_shapes_at_same_z_order(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            source = self.make_deck(directory)
            output, report = self.run_replace(directory, source)
            self.assertEqual((report["pictures_before"], report["pictures_after"]), (1, 0))
            self.assertEqual(report["replaced"][0]["native_shapes"], 13)

            slide = Presentation(output).slides[0]
            names = [shape.name for shape in slide.shapes]
            self.assertEqual(names[0], "Before")
            self.assertEqual(names[-1], "After")
            self.assertEqual(len([name for name in names if name.startswith("DIMER_ICON_ROBOT_")]), 13)
            self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes))
            self.assertTrue(all(ref.get("idx") == "0" for shape in list(slide.shapes)[1:-1] for ref in shape._element.xpath("./p:style/a:effectRef")))
            with zipfile.ZipFile(output) as deck_zip:
                xml = deck_zip.read("ppt/slides/slide1.xml")
            self.assertNotIn(b"<p:pic>", xml)

    def test_cropped_picture_is_preserved_and_reported(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            source = self.make_deck(directory, crop=True)
            output, report = self.run_replace(directory, source)
            self.assertEqual((report["pictures_before"], report["pictures_after"]), (1, 1))
            self.assertEqual(report["skipped"], [{"slide": 1, "picture": "Robot Pic", "reason": "cropped_picture"}])
            self.assertEqual(len(report["replaced"]), 0)
            self.assertTrue(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in Presentation(output).slides[0].shapes))


if __name__ == "__main__":
    unittest.main()
