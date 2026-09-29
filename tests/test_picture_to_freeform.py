from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "picture_to_freeform.py"
SPEC = importlib.util.spec_from_file_location("picture_to_freeform", SCRIPT)
assert SPEC and SPEC.loader
converter = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = converter
SPEC.loader.exec_module(converter)


def make_donut(path: Path) -> None:
    image = Image.new("RGBA", (80, 80), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse((4, 4, 76, 76), fill=(20, 100, 220, 255))
    draw.ellipse((27, 27, 53, 53), fill=(0, 0, 0, 0))
    image.save(path)


def make_two_color(path: Path) -> None:
    image = Image.new("RGBA", (80, 60), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((4, 8, 34, 52), radius=5, fill=(230, 80, 40, 255))
    draw.polygon(((47, 52), (63, 8), (76, 52)), fill=(30, 165, 90, 255))
    image.save(path)


class PictureToFreeformTests(unittest.TestCase):
    def test_converts_flat_icons_to_custom_geometry_and_preserves_z_order(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            donut = temp / "donut.png"
            two_color = temp / "two_color.png"
            source = temp / "source.pptx"
            output = temp / "output.pptx"
            make_donut(donut)
            make_two_color(two_color)

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            before = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.1), Inches(0.1), Inches(0.2), Inches(0.2))
            before.name = "before"
            first = slide.shapes.add_picture(str(donut), Inches(1), Inches(1), width=Inches(1.5), height=Inches(1.5))
            first.name = "donut"
            first.crop_left = 0.05
            first.crop_top = 0.05
            second = slide.shapes.add_picture(str(two_color), Inches(3), Inches(1), width=Inches(1.5), height=Inches(1.125))
            second.name = "two-color"
            after = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(0.1), Inches(0.2), Inches(0.2))
            after.name = "after"
            presentation.save(source)

            self.assertEqual(
                converter.main([str(source), str(output), "--picture-indices", "1,2"]),
                0,
            )

            report = json.loads(output.with_suffix(".pptx.json").read_text(encoding="utf-8"))
            self.assertEqual(
                report["summary"],
                {"converted": 2, "native_shape_count": 3, "native_node_count": 108, "pictures_before": 2, "pictures_after": 0},
            )
            self.assertEqual(report["pictures"][0]["native_shape_count"], 1)
            self.assertEqual(report["pictures"][0]["cropped_source_size_px"], {"width": 76, "height": 76})
            self.assertEqual(report["pictures"][1]["native_shape_count"], 2)

            converted = Presentation(output)
            shapes = converted.slides[0].shapes
            self.assertEqual([shape.name for shape in shapes], ["before", "DIMER_ICON_1_C1", "DIMER_ICON_2_C1", "DIMER_ICON_2_C2", "after"])
            self.assertTrue(all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in shapes))
            self.assertTrue(all(shape.shape_type == MSO_SHAPE_TYPE.FREEFORM for shape in list(shapes)[1:-1]))
            self.assertTrue(all(ref.get("idx") == "0" for shape in list(shapes)[1:-1] for ref in shape._element.xpath("./p:style/a:effectRef")))
            self.assertEqual((shapes[1].left, shapes[1].top, shapes[1].width, shapes[1].height), (first.left, first.top, first.width, first.height))

            with zipfile.ZipFile(output) as package:
                xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertNotIn("<p:pic", xml)
            self.assertGreaterEqual(xml.count("<a:custGeom"), 3)
            # The donut's one native shape has an outer contour and an inner hole.
            self.assertGreaterEqual(xml.count("<a:moveTo>"), 4)

    def test_skips_image_when_shape_budget_is_exceeded(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source_image = temp / "two_color.png"
            source = temp / "source.pptx"
            output = temp / "output.pptx"
            make_two_color(source_image)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(str(source_image), Inches(1), Inches(1))
            presentation.save(source)

            self.assertEqual(
                converter.main([str(source), str(output), "--all-pictures", "--max-shapes", "1"]),
                0,
            )
            report = json.loads(output.with_suffix(".pptx.json").read_text(encoding="utf-8"))
            self.assertEqual(report["pictures"][0]["status"], "skipped")
            self.assertEqual(report["pictures"][0]["reason"], "shape_budget_exceeded:2>1")
            self.assertEqual(report["summary"], {"skipped": 1, "pictures_before": 1, "pictures_after": 1})
            self.assertEqual(Presentation(output).slides[0].shapes[0].shape_type, MSO_SHAPE_TYPE.PICTURE)


if __name__ == "__main__":
    unittest.main()
